# -*- coding: utf-8 -*-
"""生成论文插图（matplotlib, 中文字体）"""
import os
import sys
import numpy as np

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'methodB_spectral'))

from hb_core import load_att1, load_att2, props_q1, props_q23, props_q4, C0, R0  # noqa: E402

for f in (r'C:\Windows\Fonts\msyh.ttc', r'C:\Windows\Fonts\simhei.ttf',
          r'C:\Windows\Fonts\simsun.ttc'):
    if os.path.exists(f):
        font_manager.fontManager.addfont(f)
plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'SimSun']
plt.rcParams['axes.unicode_minus'] = False
plt.rcParams['figure.dpi'] = 160
plt.rcParams['savefig.bbox'] = 'tight'


def fig_env():
    t, T, C = load_att1()
    fig, ax = plt.subplots(1, 2, figsize=(7.6, 2.5))
    ax[0].plot(t / 3600, T, lw=1.6, color='#c0392b')
    ax[0].set_xlabel('时间 / h'); ax[0].set_ylabel('烘房温度 / ℃'); ax[0].grid(alpha=.3)
    ax[0].set_title('(a) 烘房温度 $T_\\infty(t)$', fontsize=10)
    ax[1].plot(t / 3600, C, lw=1.6, color='#1f6fb4')
    ax[1].set_xlabel('时间 / h'); ax[1].set_ylabel('烘房水分浓度 / (kg/kg)'); ax[1].grid(alpha=.3)
    ax[1].set_title('(b) 烘房水分浓度 $C_\\infty(t)$', fontsize=10)
    fig.tight_layout()
    fig.savefig(os.path.join(HERE, 'fig1_env.png')); plt.close(fig)


def fig_D():
    C = np.linspace(0.03, 2.6, 800)
    T = 323.15
    d1 = 7e-9 * np.exp(-0.89 / C)
    d1w = 7e-9 * np.exp(-0.89 * C)
    d3 = 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850 / T)
    d3w = 2.4e-3 * np.exp(-0.45 * C) * np.exp(-3850 / T)
    d4 = 4.2e-4 * np.exp(-0.30 / C) * np.exp(-3850 / T)
    d4w = 4.2e-4 * np.exp(-0.30 * C) * np.exp(-3850 / T)
    fig, ax = plt.subplots(1, 2, figsize=(7.6, 2.7))
    ax[0].semilogy(C, d1, lw=1.8, label=r'正确: $7\times10^{-9}e^{-0.89/C}$')
    ax[0].semilogy(C, d1w, lw=1.8, ls='--', label=r'误读: $7\times10^{-9}e^{-0.89C}$')
    ax[0].set_xlabel('干基含水率 C / (kg/kg)'); ax[0].set_ylabel('D / (m²/s)')
    ax[0].set_title('(a) 问题1 扩散系数', fontsize=10); ax[0].legend(fontsize=7.5); ax[0].grid(alpha=.3)
    ax[1].semilogy(C, d3, lw=1.8, label=r'正确: $2.4\times10^{-3}e^{-0.45/C}e^{-3850/T}$')
    ax[1].semilogy(C, d3w, lw=1.8, ls='--', label=r'误读: $e^{-0.45C}e^{-3850/T}$')
    ax[1].semilogy(C, d4, lw=1.2, color='g', label=r'问题4: $4.2\times10^{-4}e^{-0.30/C}$')
    ax[1].set_xlabel('干基含水率 C / (kg/kg)'); ax[1].set_ylabel('D / (m²/s)')
    ax[1].set_title('(b) 问题2–4 扩散系数（T=50 ℃）', fontsize=10)
    ax[1].legend(fontsize=7.5); ax[1].grid(alpha=.3)
    fig.tight_layout()
    fig.savefig(os.path.join(HERE, 'fig2_D.png')); plt.close(fig)


def fig_profiles():
    import openpyxl
    res = os.path.join(os.path.dirname(HERE), 'methodB_spectral', 'results')
    fig, ax = plt.subplots(1, 3, figsize=(7.6, 2.5))
    # (a) 问题1 水分剖面演化
    wb = openpyxl.load_workbook(os.path.join(res, 'result1.xlsx'), data_only=True)
    rows = list(wb['水分浓度'].iter_rows(values_only=True))
    rr = np.array(rows[0][1:], float)
    for k, lab in [(300, '300 s'), (900, '900 s'), (1800, '1800 s')]:
        ax[0].plot(rr, rows[k + 1][1:], lw=1.4, label=lab)
    ax[0].set_xlabel('到中心距离 / cm'); ax[0].set_ylabel('C / (kg/kg)')
    ax[0].set_title('(a) 问题1 水分剖面', fontsize=10); ax[0].legend(fontsize=8); ax[0].grid(alpha=.3)
    # (b) 问题3
    wb = openpyxl.load_workbook(os.path.join(res, 'result3.xlsx'), data_only=True)
    rows = list(wb['水分浓度'].iter_rows(values_only=True))
    tt = np.array([r[0] for r in rows[1:]], float)
    M = np.array([r[1:] for r in rows[1:]], float)
    for hh in (6, 18, 36, 54):
        k = int(np.argmin(np.abs(tt - hh * 3600)))
        ax[1].plot(rr, M[k], lw=1.4, label='%d h' % hh)
    ax[1].plot(rr, M[-1], lw=1.6, color='k', label='结束 %.2f h' % (tt[-1] / 3600))
    ax[1].axhline(0.15, color='r', ls=':', lw=1)
    ax[1].set_xlabel('到中心距离 / cm'); ax[1].set_ylabel('C / (kg/kg)')
    ax[1].set_title('(b) 问题3 水分剖面', fontsize=10); ax[1].legend(fontsize=8); ax[1].grid(alpha=.3)
    # (c) 问题4 中心与表面
    wb = openpyxl.load_workbook(os.path.join(res, 'result4.xlsx'), data_only=True)
    rows = list(wb['水分浓度'].iter_rows(values_only=True))
    tt4 = np.array([r[0] for r in rows[1:]], float)
    M4 = np.array([[np.nan if v is None else v for v in r[1:]] for r in rows[1:]], float)
    ax[2].plot(tt4 / 3600, M4[:, 0], lw=1.6, label='中心 (0 cm)')
    ax[2].plot(tt4 / 3600, M4[:, -1], lw=1.6, label='表面')
    ax[2].axhline(0.15, color='r', ls=':', lw=1, label='判据 0.15')
    ax[2].set_xlabel('时间 / h'); ax[2].set_ylabel('C / (kg/kg)')
    ax[2].set_title('(c) 问题4 中心/表面', fontsize=10); ax[2].legend(fontsize=8); ax[2].grid(alpha=.3)
    fig.tight_layout()
    fig.savefig(os.path.join(HERE, 'fig3_profiles.png')); plt.close(fig)


def fig_verify():
    fig, ax = plt.subplots(1, 2, figsize=(7.6, 2.6))
    n = np.array([43, 55, 57, 81])
    tf = np.array([205809.6, 205809.8, 205809.7, 205809.8])
    ax[0].plot(n, tf - 205809.7, 'o-', lw=1.4)
    ax[0].axhspan(-1, 1, color='g', alpha=.12)
    ax[0].set_xlabel('自由度 ndof'); ax[0].set_ylabel('$t_{dry}$ 偏差 / s')
    ax[0].set_title('(a) 问题3 烘干时长的网格收敛', fontsize=10); ax[0].grid(alpha=.3)
    lab = ['表1 温度', '表1 水分', '表3 温度', '表3 水分', '表5 水分', '表6 水分']
    d = [0.0019, 1.2e-4, 0.0024, 5.5e-5, 7.9e-5, 5.8e-5]
    ax[1].barh(lab, d, color='#1f6fb4')
    ax[1].axvline(1e-4, color='r', ls='--', lw=1.2, label='四位小数分辨力 $10^{-4}$')
    ax[1].set_xscale('log'); ax[1].set_xlabel('与工作区实现的逐格最大差')
    ax[1].set_title('(b) 跨方法一致性', fontsize=10); ax[1].legend(fontsize=8); ax[1].grid(alpha=.3, axis='x')
    fig.tight_layout()
    fig.savefig(os.path.join(HERE, 'fig4_verify.png')); plt.close(fig)


if __name__ == '__main__':
    fig_env(); fig_D(); fig_profiles(); fig_verify()
    print('figures ok:', [f for f in os.listdir(HERE) if f.endswith('.png')])
