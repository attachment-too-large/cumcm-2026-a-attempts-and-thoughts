# -*- coding: utf-8 -*-
"""独立数值复核（不使用题库代码）：均匀 xi 网格节点中心 FV + 调和平均面系数 + 隐式欧拉。"""
import sys, io, os, time
if sys.stdout.encoding is None or 'utf' not in (sys.stdout.encoding or '').lower():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import numpy as np
import openpyxl
from scipy.interpolate import PchipInterpolator
from scipy.linalg import solve_banded

ROOT = r'C:\Users\qing1\Desktop\A题'
ATT = os.path.join(ROOT, '附件')

wb = openpyxl.load_workbook(os.path.join(ATT, '附件1.xlsx'), data_only=True)
rows = [r for r in wb['Sheet1'].iter_rows(values_only=True)][1:]
t_env = np.array([float(r[0]) for r in rows])
T_env = np.array([float(r[1]) for r in rows])
C_env = np.array([float(r[2]) for r in rows])
fT = PchipInterpolator(t_env, T_env)
fC = PchipInterpolator(t_env, C_env)
T_hold, C_hold = float(fT(t_env[-1])), float(fC(t_env[-1]))
t_hold = float(t_env[-1])

wb = openpyxl.load_workbook(os.path.join(ATT, '附件2.xlsx'), data_only=True)
rows = [r for r in wb['Sheet1'].iter_rows(values_only=True)][1:]
t_R = np.array([float(r[0]) for r in rows])
R_cm = np.array([float(r[1]) for r in rows])
fR = PchipInterpolator(t_R, R_cm)

h_conv, km = 25.0, 8e-7
T0, C0 = 28.0, 2.55


def Tinf(t):
    return float(fT(t)) if t <= t_hold else T_hold


def Cinf(t):
    return float(fC(t)) if t <= t_hold else C_hold


def Rof(t):
    return (float(fR(t)) if t <= t_R[-1] else float(R_cm[-1])) * 1e-2


def props(mode, C, T_C):
    C = np.maximum(C, 1e-12)
    TK = T_C + 273.15
    if mode == 'p1':
        # 附录 2（问题 1）：常物性
        rho = np.full_like(C, 820.0)
        cp = np.full_like(C, 2600.0)
        k = np.full_like(C, 0.36)
        D = 7e-9 * np.exp(-0.89 / C)
    elif mode == 'p3':
        rho = 650.0 + 128.0 * C
        cp = 1450.0 + 2736.0 * C / (C + 1.0)
        k = 0.21 + 0.38 * C / (C + 1.0)
        D = 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / TK)
    elif mode == 'p4':
        rho = 760.0 + 90.0 * C
        cp = 1850.0 + 2150.0 * C / (C + 1.0)
        k = 0.12 + 0.20 * C / (C + 1.0)
        D = 4.2e-4 * np.exp(-0.30 / C) * np.exp(-3850.0 / TK)
    else:
        raise ValueError(mode)
    return rho, cp, k, D


class FV:
    def __init__(self, mode, shrink, N=400):
        self.mode, self.shrink, self.N = mode, shrink, N
        xi = np.linspace(0.0, 1.0, N + 1)
        self.xi = 0.5 * (xi[:-1] + xi[1:])                 # 节点
        self.xf = xi[1:-1]                                  # 内面
        lo = np.concatenate([[0.0], self.xf])
        hi = np.concatenate([self.xf, [1.0]])
        self.wV = 0.5 * (hi ** 2 - lo ** 2)                 # 体积权重（*R^2）
        self.Ain = lo                                       # 内侧面权重（*R）
        self.Aout = hi                                      # 外侧面权重（*R）
        self.dface = np.diff(self.xi)                       # 相邻节点距离（相对）
        self.Af = self.xf                                   # 内面面积权重（*R）

    def step(self, T, C, t, dt):
        N = self.N
        R = Rof(t + 0.5 * dt) if self.shrink else 0.02
        V = R ** 2 * self.wV
        Ti, Ci = Tinf(t), Cinf(t)
        rho, cp, k, D = props(self.mode, C, T)
        kf = 2 * k[:-1] * k[1:] / (k[:-1] + k[1:] + 1e-300)
        Df = 2 * D[:-1] * D[1:] / (D[:-1] + D[1:] + 1e-300)
        Gf_T = R * self.Af * kf / (R * self.dface)
        Gf_C = R * self.Af * Df / (R * self.dface)
        capT = rho * cp * V
        T = self._solve(T, capT, Gf_T, R * 1.0 * h_conv, Ti, dt)
        C = self._solve(C, V.copy(), Gf_C, R * 1.0 * km, Ci, dt)
        return T, C

    def _solve(self, Y, Cap, Gf, Gout, Yinf, dt):
        N = self.N
        g = np.concatenate([Gf, [Gout]])                   # 上侧面通量系数（长度 N）
        lo = np.concatenate([[0.0], Gf])                    # 下侧面（长度 N）
        diag = Cap / dt + lo + g
        ab = np.zeros((3, N))
        ab[0, 1:] = -g[:-1]
        ab[1, :] = diag
        ab[2, :-1] = -lo[1:]
        rhs = Cap / dt * Y
        rhs[N - 1] += Gout * Yinf
        return solve_banded((1, 1), ab, rhs)

    def till(self, thr, t_end, dt=60.0, trace=None):
        N = self.N
        T = np.full(N, T0); C = np.full(N, C0); t = 0.0
        while t < t_end - 1e-12:
            T, C = self.step(T, C, t, dt)
            t += dt
            if trace and abs(t % trace) < 1e-9:
                print(f'      t={t/3600:8.3f} h C_c={C[0]:.6f} C_s={C[-1]:.4e} '
                      f'T_c={T[0]:.3f} T_s={T[-1]:.3f}', flush=True)
            if C.max() < thr:
                return t, C, T
        return None, C, T


if __name__ == '__main__':
    thr = 0.15
    print('=' * 74)
    print('独立复核 A：问题 3（附录3，R 固定 2 cm，km=8e-7, h=25）')
    print('=' * 74)
    res = {}
    for N in (100, 200, 400, 800):
        t0 = time.time()
        s = FV('p3', False, N=N)
        tt, C, T = s.till(thr, 3600 * 90, 60.0)
        res[N] = tt / 3600
        print(f'  N={N:4d}: 首次 max(C)<0.15 于 {tt:.0f}s = {tt/3600:.4f} h '
              f'(C_max={C.max():.3e}, C_center={C[0]:.6f}, T_center={T[0]:.3f}) '
              f'[{time.time()-t0:.0f}s]', flush=True)
    print('  → 网格收敛序列:', {k: round(v, 4) for k, v in res.items()})
    print()
    print('=' * 74)
    print('独立复核 B：问题 4（附录4，仿射收缩 R(t)）')
    print('=' * 74)
    res4 = {}
    for N in (200, 400, 800):
        t0 = time.time()
        s = FV('p4', True, N=N)
        tt, C, T = s.till(thr, 3600 * 80, 60.0)
        res4[N] = tt / 3600
        print(f'  N={N:4d}: 首次 max(C)<0.15 于 {tt:.0f}s = {tt/3600:.4f} h '
              f'(C_max={C.max():.3e}, C_center={C[0]:.6f}) [{time.time()-t0:.0f}s]',
              flush=True)
    print('  → 网格收敛序列:', {k: round(v, 4) for k, v in res4.items()})
