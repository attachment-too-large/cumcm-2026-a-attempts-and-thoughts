# -*- coding: utf-8 -*-
"""审计脚本 1：附件数据 + 结果文件模板/输出 的独立核查。"""
import sys, io, os, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import numpy as np
import openpyxl

ROOT = r'C:\Users\qing1\Desktop\A题'
ATT = os.path.join(ROOT, '附件')
OUT = os.path.join(ROOT, 'cumcm2026-a-herb-drying-main', 'outputs', 'result_files')


def sheet_info(path, name=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[0] if rows else ()
        times = [r[0] for r in rows[1:] if r and r[0] is not None]
        out[ws.title] = dict(n_rows=len(rows), n_cols=ws.max_column, header=hdr,
                             t_first=times[0] if times else None,
                             t_last=times[-1] if times else None,
                             n_times=len(times))
    wb.close()
    return out


def fmt(v):
    return v


print('#' * 78)
print('# 1. 附件1 / 附件2 原始数据')
print('#' * 78)
wb = openpyxl.load_workbook(os.path.join(ATT, '附件1.xlsx'), data_only=True)
for ws in wb.worksheets:
    rows = [r for r in ws.iter_rows(values_only=True)]
    print(f'-- 附件1 sheet="{ws.title}" rows={len(rows)} cols={ws.max_column}')
    print('   header:', rows[0])
    print('   first3:', rows[1:4])
    print('   last3 :', rows[-3:])
    t = np.array([float(r[0]) for r in rows[1:] if r[0] is not None])
    T = np.array([float(r[1]) for r in rows[1:] if r[0] is not None])
    C = np.array([float(r[2]) for r in rows[1:] if r[0] is not None])
    print(f'   t: {t[0]}..{t[-1]}  n={len(t)}  dt uniq={np.unique(np.diff(t))[:5]}')
    print(f'   T: min={T.min():.4f} max={T.max():.4f} first={T[0]:.4f} last={T[-1]:.4f}')
    print(f'   C: min={C.min():.4f} max={C.max():.4f} first={C[0]:.4f} last={C[-1]:.4f}')
    print(f'   T last-10: {np.round(T[-10:],3)}')
    print(f'   C last-10: {np.round(C[-10:],4)}')
    # 平台段统计（后半段）
    half = len(t)//2
    print(f'   T 平台段(后半) mean={T[half:].mean():.4f} std={T[half:].std():.4f} '
          f'min={T[half:].min():.4f} max={T[half:].max():.4f}')
    print(f'   C 平台段(后半) mean={C[half:].mean():.4f} std={C[half:].std():.4f} '
          f'min={C[half:].min():.4f} max={C[half:].max():.4f}')
    # 单调性
    print(f'   T monotone increasing? {bool(np.all(np.diff(T) >= -1e-12))}; '
          f'最大回退 {float(np.min(np.diff(T))):.4f}')
    print(f'   C monotone? {bool(np.all(np.diff(C) >= -1e-12))}; 最大回退 {float(np.min(np.diff(C))):.4f}')

wb = openpyxl.load_workbook(os.path.join(ATT, '附件2.xlsx'), data_only=True)
for ws in wb.worksheets:
    rows = [r for r in ws.iter_rows(values_only=True)]
    print(f'-- 附件2 sheet="{ws.title}" rows={len(rows)} cols={ws.max_column}')
    print('   header:', rows[0])
    print('   first5:', rows[1:6])
    print('   last3 :', rows[-3:])
    t = np.array([float(r[0]) for r in rows[1:] if r[0] is not None])
    R = np.array([float(r[1]) for r in rows[1:] if r[0] is not None])
    print(f'   t: {t[0]}..{t[-1]} n={len(t)} dt uniq={np.unique(np.diff(t))[:5]}')
    print(f'   R: first={R[0]} min={R.min()} last={R[-1]}')
    print(f'   R monotone non-increasing? {bool(np.all(np.diff(R) <= 1e-12))}; 最大回升 {float(np.max(np.diff(R))):.6f}')
    print('   R 前20:', R[:20])
    print('   R 后10:', R[-10:])
    # 收缩速率
    dR = np.diff(R)/np.diff(t)
    i = int(np.argmax(np.abs(dR)))
    print(f'   |dR/dt| max = {abs(dR[i]):.3e} cm/s at t={t[i]}s')

print()
print('#' * 78)
print('# 2. 结果文件模板（附件3）')
print('#' * 78)
for f in ['result1.xlsx', 'result2.xlsx', 'result3.xlsx', 'result4.xlsx']:
    info = sheet_info(os.path.join(ATT, '附件3', f))
    print(f'== {f}')
    for k, v in info.items():
        print(f'   [{k}] rows={v["n_rows"]} cols={v["n_cols"]} header={v["header"]}'
              f' tf={v["t_first"]} tl={v["t_last"]}')

print()
print('#' * 78)
print('# 3. 本文输出的结果文件')
print('#' * 78)
for f in sorted(os.listdir(OUT)):
    if not f.endswith('.xlsx'):
        continue
    p = os.path.join(OUT, f)
    print(f'== {f}  ({os.path.getsize(p)/1024/1024:.2f} MB)')
    info = sheet_info(p)
    for k, v in info.items():
        print(f'   [{k}] rows={v["n_rows"]} cols={v["n_cols"]} tf={v["t_first"]} tl={v["t_last"]}')
        print(f'        header={v["header"]}')
