# -*- coding: utf-8 -*-
"""生成 result1-4.xlsx（按题目模板）。

约定：
  * A 列 = 时间（s）；第 1 行为到药材中心的距离（0, 0.1, ..., 2.0 cm，共 21 列）。
  * result4 的列名为"当前物理距离"，r_j > R(t_i) 处该位置不存在（留空），
    并在最后附加"药材表面"列给出当前表面值 C(R(t),t)。
  * 所有数值保留 4 位小数。
输出目录：<包根>/结果/
"""
import os
from console_utf8 import fix_console
fix_console()
import numpy as np
import openpyxl
import common as cm
from spectral import HerbSolver

OUTDIR = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                      "..", "结果"))
os.makedirs(OUTDIR, exist_ok=True)

RCOLS = np.round(np.arange(0.0, 2.0001, 0.1), 10)      # 0.0 .. 2.0 cm
RC = RCOLS * 1e-2                                       # m
CRIT = 0.15
N = 64


def write_sheet(ws, times, values, header_last=None):
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, rc in enumerate(RCOLS):
        ws.cell(row=1, column=2 + j, value=float(rc))
    if header_last is not None:
        ws.cell(row=1, column=2 + len(RCOLS), value=header_last)
    for i in range(len(times)):
        ws.cell(row=2 + i, column=1, value=float(times[i]))
        for j in range(len(RCOLS)):
            v = values[i, j]
            if v is not None and np.isfinite(v):
                ws.cell(row=2 + i, column=2 + j, value=round(float(v), 4))
    if header_last is not None:
        for i in range(len(times)):
            v = values[i, len(RCOLS)]
            if v is not None and np.isfinite(v):
                ws.cell(row=2 + i, column=2 + len(RCOLS), value=round(float(v), 4))


def sample(s, sol, u_targets, which):
    out = np.empty((sol.y.shape[1], len(u_targets)))
    for i in range(sol.y.shape[1]):
        C, T = s.interp_u(sol.y[:, i], u_targets)
        out[i] = T if which == "T" else C
    return out


# ================= result1 =================
print("result1 ...")
s = HerbSolver(N, prob=1)
tt = np.arange(1.0, 1801.0, 1.0)
at = np.concatenate([np.full(s.Np, 1e-13), np.full(s.Np, 1e-11)])
sol = s.solve(1800.0, t_eval=tt, rtol=1e-11, atol=at)
U = (RC / cm.R0) ** 2
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
write_sheet(ws, tt, sample(s, sol, U, "T"))
ws2 = wb.create_sheet("水分浓度")
write_sheet(ws2, tt, sample(s, sol, U, "C"))
wb.save(os.path.join(OUTDIR, "result1.xlsx"))

# ================= result2 =================
print("result2 ...")
s2 = HerbSolver(N, prob=2)
tt2 = np.arange(1.0, 10801.0, 1.0)
at2 = np.concatenate([np.full(s2.Np, 1e-12), np.full(s2.Np, 1e-10)])
sol2 = s2.solve(10800.0, t_eval=tt2, rtol=1e-11, atol=at2)
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
write_sheet(ws, tt2, sample(s2, sol2, U, "T"))
ws2 = wb.create_sheet("水分浓度")
write_sheet(ws2, tt2, sample(s2, sol2, U, "C"))
wb.save(os.path.join(OUTDIR, "result2.xlsx"))


def dry(s, shrink=False, rtol=1e-11):
    at = np.concatenate([np.full(s.Np, 1e-13), np.full(s.Np, 1e-11)])

    def ev(t, y):
        return np.max(y[:s.Np]) - CRIT
    ev.terminal = True
    ev.direction = -1
    e = s.solve(300000.0, rtol=rtol, atol=at, events=ev)
    return float(e.t_events[0][0])


# ================= result3 =================
print("result3 ...")
s3 = HerbSolver(N, prob=3)
td3 = dry(s3)
tt3 = np.arange(60.0, np.floor(td3 / 60.0) * 60.0 + 1.0, 60.0)
tt3 = np.append(tt3, td3)
at3 = np.concatenate([np.full(s3.Np, 1e-13), np.full(s3.Np, 1e-11)])
sol3 = s3.solve(td3, t_eval=list(tt3), rtol=1e-11, atol=at3)
U3 = (RC / cm.R0) ** 2
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
write_sheet(ws, tt3, sample(s3, sol3, U3, "C"))
wb.save(os.path.join(OUTDIR, "result3.xlsx"))

# ================= result4 =================
print("result4 ...")
s4 = HerbSolver(N, prob=4, shrink=True)
td4 = dry(s4, shrink=True)
tt4 = np.arange(60.0, np.floor(td4 / 60.0) * 60.0 + 1.0, 60.0)
tt4 = np.append(tt4, td4)
at4 = np.concatenate([np.full(s4.Np, 1e-13), np.full(s4.Np, 1e-11)])
sol4 = s4.solve(td4, t_eval=list(tt4), rtol=1e-11, atol=at4)
vals = np.full((len(tt4), len(RCOLS) + 1), np.nan)
for i, t in enumerate(tt4):
    Rt = s4.R_of(t)
    for j, r in enumerate(RC):
        if r <= Rt + 1e-12:                 # 该物理位置仍然存在
            Cn, _ = s4.interp_u(sol4.y[:, i], (r / Rt) ** 2)
            vals[i, j] = Cn[0]
    vals[i, -1] = sol4.y[s4.N, i]           # 当前表面值
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
write_sheet(ws, tt4, vals, header_last="药材表面")
wb.save(os.path.join(OUTDIR, "result4.xlsx"))

print("result1: %d 行 x %d 列" % (len(tt), len(RCOLS) + 1))
print("result2: %d 行 x %d 列" % (len(tt2), len(RCOLS) + 1))
print("result3: %d 行, tdry=%.3f s" % (len(tt3), td3))
print("result4: %d 行, tdry=%.3f s" % (len(tt4), td4))
print("已写入:", OUTDIR)

