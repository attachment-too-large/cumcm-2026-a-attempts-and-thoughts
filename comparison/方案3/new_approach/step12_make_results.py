# -*- coding: utf-8 -*-
"""生成 result1-4.xlsx（新解法结果）。

约定：
  * A 列 = 时间（s）；第 1 行为到药材中心的距离（0, 0.1, ..., 2.0 cm，共 21 列）。
  * result4 的列名为"当前物理距离"，r_j > R(t_i) 处该位置不存在（留空），
    并在最后附加"药材表面"列给出当前表面值 C(R(t),t)。
  * 所有数值保留 4 位小数。
"""
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter
import common as cm
from spectral import HerbSolver

RCOLS = np.round(np.arange(0.0, 2.0001, 0.1), 10)      # 0.0 .. 2.0 cm
RC = RCOLS * 1e-2                                       # m
CRIT = 0.15


def write_sheet(ws, times, values, header_last=None):
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, rc in enumerate(RCOLS):
        ws.cell(row=1, column=2 + j, value=float(rc))
    if header_last is not None:
        ws.cell(row=1, column=2 + len(RCOLS), value=header_last)
    for i, t in enumerate(times):
        ws.cell(row=2 + i, column=1, value=float(t))
        for j in range(len(RCOLS)):
            v = values[i, j]
            if v is not None and np.isfinite(v):
                ws.cell(row=2 + i, column=2 + j, value=round(float(v), 4))
    if header_last is not None:
        for i in range(len(times)):
            v = values[i, len(RCOLS)]
            if v is not None and np.isfinite(v):
                ws.cell(row=2 + i, column=2 + len(RCOLS), value=round(float(v), 4))


def sample_table(s, sol, u_targets):
    n = sol.y.shape[1]
    out = np.empty((n, len(u_targets)))
    for i in range(n):
        C, T = s.interp_u(sol.y[:, i], u_targets)
        out[i] = C
    return out


def sample_T(s, sol, u_targets):
    n = sol.y.shape[1]
    out = np.empty((n, len(u_targets)))
    for i in range(n):
        C, T = s.interp_u(sol.y[:, i], u_targets)
        out[i] = T
    return out


# ================= result1 =================
print("result1 ...")
N = 64
s = HerbSolver(N, prob=1)
tt = np.arange(1.0, 1801.0, 1.0)
at = np.concatenate([np.full(s.Np, 1e-13), np.full(s.Np, 1e-11)])
sol = s.solve(1800.0, t_eval=tt, rtol=1e-11, atol=at)
U = (RC / cm.R0) ** 2
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
write_sheet(ws, tt, sample_T(s, sol, U))
ws2 = wb.create_sheet("水分浓度")
write_sheet(ws2, tt, sample_table(s, sol, U))
wb.save(r"C:\Users\qing1\Desktop\数A\附件\附件3\result1.xlsx")

# ================= result2 =================
print("result2 ...")
s2 = HerbSolver(N, prob=2)
tt2 = np.arange(1.0, 10801.0, 1.0)
at2 = np.concatenate([np.full(s2.Np, 1e-12), np.full(s2.Np, 1e-10)])
sol2 = s2.solve(10800.0, t_eval=tt2, rtol=1e-11, atol=at2)
U2 = (RC / cm.R0) ** 2
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
write_sheet(ws, tt2, sample_T(s2, sol2, U2))
ws2 = wb.create_sheet("水分浓度")
write_sheet(ws2, tt2, sample_table(s2, sol2, U2))
wb.save(r"C:\Users\qing1\Desktop\数A\附件\附件3\result2.xlsx")

# ================= result3 =================
print("result3 ...")
s3 = HerbSolver(N, prob=3)
at3 = np.concatenate([np.full(s3.Np, 1e-13), np.full(s3.Np, 1e-11)])


def ev3(t, y):
    return np.max(y[:s3.Np]) - CRIT
ev3.terminal = True; ev3.direction = -1
solE = s3.solve(300000.0, rtol=1e-11, atol=at3, events=ev3)
tdry3 = float(solE.t_events[0][0])
tt3 = np.arange(60.0, np.floor(tdry3 / 60.0) * 60.0 + 1.0, 60.0)
tt3 = np.append(tt3, tdry3)
sol3 = s3.solve(tdry3, t_eval=tt3, rtol=1e-11, atol=at3)
U3 = (RC / cm.R0) ** 2
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
write_sheet(ws, tt3, sample_table(s3, sol3, U3))
wb.save(r"C:\Users\qing1\Desktop\数A\附件\附件3\result3.xlsx")

# ================= result4 =================
print("result4 ...")
s4 = HerbSolver(N, prob=4, shrink=True)
at4 = np.concatenate([np.full(s4.Np, 1e-13), np.full(s4.Np, 1e-11)])


def ev4(t, y):
    return np.max(y[:s4.Np]) - CRIT
ev4.terminal = True; ev4.direction = -1
solE4 = s4.solve(300000.0, rtol=1e-11, atol=at4, events=ev4)
tdry4 = float(solE4.t_events[0][0])
tt4 = np.arange(60.0, np.floor(tdry4 / 60.0) * 60.0 + 1.0, 60.0)
tt4 = np.append(tt4, tdry4)
sol4 = s4.solve(tdry4, t_eval=tt4, rtol=1e-11, atol=at4)
vals = np.full((len(tt4), len(RCOLS) + 1), np.nan)
for i, t in enumerate(tt4):
    Rt = s4.R_of(t)
    for j, r in enumerate(RC):
        if r <= Rt + 1e-12:
            Cn, _ = s4.interp_u(sol4.y[:, i], (r / Rt) ** 2)
            vals[i, j] = Cn[0]
    vals[i, -1] = sol4.y[s4.N, i]        # 当前表面
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
write_sheet(ws, tt4, vals, header_last="药材表面")
wb.save(r"C:\Users\qing1\Desktop\数A\附件\附件3\result4.xlsx")

rep = []
rep.append("result1: %d 行 x %d 列; result2: %d 行 x %d 列"
           % (len(tt), len(RCOLS) + 1, len(tt2), len(RCOLS) + 1))
rep.append("result3: %d 行, tdry=%.3f s" % (len(tt3), tdry3))
rep.append("result4: %d 行, tdry=%.3f s" % (len(tt4), tdry4))
open(r"C:\Users\qing1\Desktop\数A\new_approach\step12_results.txt", "w",
     encoding="utf-8").write("\n".join(rep))
print("done")
