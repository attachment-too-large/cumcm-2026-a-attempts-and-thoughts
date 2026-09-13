# -*- coding: utf-8 -*-
# 临时检查脚本: 核对四个结果 xlsx 的结构与附件 3 模板的一致性。
import os
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for f in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"):
    p = os.path.join(BASE, "results", f)
    wb = openpyxl.load_workbook(p, read_only=True)
    print("==", f, wb.sheetnames)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
        print("   sheet:", ws.title, " dims:", ws.max_row, "x", ws.max_column)
        print("   header[:6] =", rows[0][:6])
        print("   header[-2:] =", rows[0][-2:])
        print("   row1[:6] =", rows[1][:6])
        print("   row1[-2:] =", rows[1][-2:])
        print("   last row =", list(ws.iter_rows(min_row=ws.max_row,
                                                max_row=ws.max_row,
                                                values_only=True))[0][:6])
    wb.close()
