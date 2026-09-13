# -*- coding: utf-8 -*-
"""方法B 生产计算：问题1-4

建模约定（本方法独立判断）：
 * 空间：材料坐标 sigma=r/R(t)；分域谱元；弱形式，表面 Robin 自然边界
 * 环境：附件1 的 T_inf(t), C_inf(t)，0~14400 s 实测曲线线性插值；
          t>14400 s 保持末端值（=恒温干燥段设定值 50.165 ℃ / 0.04986）
 * 问题1 用附录2 参数；问题2/3 用附录3；问题4 用附录4 + 附件2 的 R(t)
 * 问题4 输出网格采用材料(Lagrangian)坐标 0,0.1,...,2.0 cm，末列即“药材表面”
"""
import json
import os
import time
import numpy as np

from hb_core import (Ambient, Radius, load_att1, load_att2, props_q1, props_q23,
                     props_q4, R0, T0_C, C0, H_CONV, HM_MASS, DATA_DIR)
from sem_core import SEM1D, Model, drying_time

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'results')
os.makedirs(OUT, exist_ok=True)

MESH = dict(breaks=(0.0, 0.5, 0.8, 0.95, 1.0), n=20)   # 精网格（ndof=81）
MESH_LONG = dict(breaks=(0.0, 0.6, 0.9, 1.0), n=18)     # 长时程网格（ndof=55）
#   收敛性：Q3 的 t_f 在 ndof=43/55/57 下为 205809.6/205809.8/205809.7 s（差 0.2 s）
GRID_CM = np.round(np.arange(0, 20.0001, 1.0) / 10.0, 4)


def build(props, radius=None, mesh=None, mode='linear'):
    t, T, C = load_att1()
    amb = Ambient(t, T, C, mode=mode)
    S = SEM1D(**(mesh or MESH))
    return Model(S, props, amb, radius)


PROPS_Q1 = dict(props_q1())
PROPS_Q1['const'] = True          # 常物性快速通道


def report(model, y, t):
    T, C = model.sample(y, t)
    return T, C


# ------------------------------------------------------------------ 问题1
def q1(mesh=None, mode='linear', rtol=1e-10, atol=1e-12, method='BDF'):
    print('[Q1] 预热平衡段 0-1800 s ...', flush=True)
    m = build(PROPS_Q1, None, mesh, mode)
    t0 = time.time()
    sol = m.solve(1800.0, t_eval=np.arange(0, 1800.0001, 1.0), rtol=rtol, atol=atol,
                  method=method)
    print('     %.1fs, ndof=%d, nfev=%d' % (time.time() - t0, m.ndof, sol.nfev), flush=True)
    T = np.array([m.sample(sol.y[:, k], sol.t[k])[0] for k in range(len(sol.t))])
    C = np.array([m.sample(sol.y[:, k], sol.t[k])[1] for k in range(len(sol.t))])
    tab = {}
    t, Tt, Ct = load_att1()
    for tt in (100, 300, 600, 900, 1200, 1500, 1800):
        k = int(round(tt))
        tab[tt] = dict(T=[round(float(v), 4) for v in T[k]],
                       C=[round(float(v), 4) for v in C[k]])
    return dict(model=m, sol=sol, t=sol.t, T=T, C=C, table=tab)


# ------------------------------------------------------------------ 问题2
def q2(mesh=None, mode='linear', rtol=1e-10, atol=1e-12, t_end=10800.0, method='BDF'):
    print('[Q2] 全流程 0-%.0f s ...' % t_end, flush=True)
    m = build(props_q23(), None, mesh, mode)
    t0 = time.time()
    sol = m.solve(t_end, t_eval=np.arange(0, t_end + 0.0001, 1.0), rtol=rtol, atol=atol,
                  method=method)
    print('     %.1fs, nfev=%d' % (time.time() - t0, sol.nfev), flush=True)
    T = np.array([m.sample(sol.y[:, k], sol.t[k])[0] for k in range(len(sol.t))])
    C = np.array([m.sample(sol.y[:, k], sol.t[k])[1] for k in range(len(sol.t))])
    tab = {}
    for hh in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
        k = int(round(hh * 3600))
        tab[hh] = dict(T=[round(float(v), 4) for v in T[k]],
                       C=[round(float(v), 4) for v in C[k]])
    return dict(model=m, sol=sol, t=sol.t, T=T, C=C, table=tab)


# ------------------------------------------------------------------ 问题3/4
def q34(props, radius=None, tag='Q3', mesh=None, mode='linear',
        rtol=1e-9, atol=1e-11, t_max=600000.0, step=60.0, method='BDF'):
    print('[%s] 烘干至各处 C<0.15 ...' % tag, flush=True)
    m = build(props, radius, mesh, mode)
    t0 = time.time()
    tf, sol = drying_time(m, t_max=t_max, rtol=rtol, atol=atol)
    print('     t_f = %.2f s = %.4f h   (%.1fs, nfev=%d)'
          % (tf, tf / 3600.0, time.time() - t0, sol.nfev), flush=True)
    # 输出网格
    tg = np.arange(0.0, tf + 1e-9, step)
    if tg[-1] < tf - 1e-9:
        tg = np.concatenate([tg, [tf]])
    else:
        tg[-1] = tf
    sol2 = m.solve(tf, t_eval=tg, rtol=rtol, atol=atol, method=method)
    T = np.array([m.sample(sol2.y[:, k], sol2.t[k])[0] for k in range(len(sol2.t))])
    C = np.array([m.sample(sol2.y[:, k], sol2.t[k])[1] for k in range(len(sol2.t))])
    # 单调性检查：C 应随 sigma 递减（中心最大）
    mono = bool(np.all(np.diff(C[-1]) <= 1e-12))
    print('     末态沿 sigma 单调递减: %s   C(0)=%.6f  C(R)=%.6f'
          % (mono, C[-1][0], C[-1][-1]), flush=True)
    tab = {}
    hh = 6.0
    while hh * 3600.0 < tf - 1e-9:
        k = int(np.argmin(np.abs(sol2.t - hh * 3600.0)))
        tab[hh] = dict(C=[round(float(v), 4) for v in C[k]], t=float(sol2.t[k]))
        hh += 6.0
    tab['end'] = dict(C=[round(float(v), 4) for v in C[-1]], t=float(sol2.t[-1]))
    return dict(model=m, sol=sol2, tf=tf, t=sol2.t, T=T, C=C, table=tab, mono=mono)


# ------------------------------------------------------------------ xlsx 输出
def write_xlsx(path, times, sheets, surface_last=False, surface_extra=False,
               extra_sheet=None):
    """sheets: {'温度': (Tmat,), ...}；times: (nt,)；NaN -> 空单元格
    surface_last : 21 个固定列中最后一列(2.0cm)改标签为“药材表面”（材料坐标版）
    surface_extra: 在 21 个固定列后再附加一列“药材表面”（物理坐标版）"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    base = [float(v) for v in GRID_CM]
    if surface_last:
        base[-1] = '药材表面'
    hdr = ['时间\\到药材中心的距离'] + base + (['药材表面'] if surface_extra else [])
    for name, mat in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(hdr)
        assert mat.shape[1] == len(hdr) - 1, (mat.shape, len(hdr))
        for i, tt in enumerate(times):
            row = [float(tt)]
            for v in np.round(mat[i], 4):
                row.append(None if not np.isfinite(v) else float(v))
            ws.append(row)
    if extra_sheet is not None:
        for name, (h, data) in extra_sheet.items():
            ws = wb.create_sheet(name)
            ws.append(h)
            for r in data:
                ws.append(list(r))
    wb.save(path)
    print('  -> %s (%d 行 x %d 列)' % (os.path.basename(path), len(times), len(hdr)))


if __name__ == '__main__':
    import sys
    which = sys.argv[1] if len(sys.argv) > 1 else 'all'
    res = {}
    if which in ('all', 'q1'):
        r1 = q1()
        write_xlsx(os.path.join(OUT, 'result1.xlsx'), r1['t'],
                   {'温度': r1['T'], '水分浓度': r1['C']})
        res['q1'] = r1['table']
    if which in ('all', 'q2'):
        r2 = q2()
        write_xlsx(os.path.join(OUT, 'result2.xlsx'), r2['t'],
                   {'温度': r2['T'], '水分浓度': r2['C']})
        res['q2'] = r2['table']
    if which in ('all', 'q3'):
        r3 = q34(props_q23(), None, 'Q3', mesh=MESH_LONG)
        write_xlsx(os.path.join(OUT, 'result3.xlsx'), r3['t'], {'水分浓度': r3['C']})
        res['q3'] = dict(tf_s=r3['tf'], tf_h=r3['tf'] / 3600.0, table=r3['table'],
                         mono=r3['mono'])
    if which in ('all', 'q4'):
        t2, R2 = load_att2()
        rad = Radius(fixed=False, t=t2, R=R2)
        r4 = q34(props_q4(), rad, 'Q4', mesh=MESH_LONG)
        m4 = r4['model']
        # (a) 物理(当前)构形坐标：固定 0-2.0 cm 列，材料外留空 + 表面列 + 表面位置表
        P4 = np.array([m4.sample(r4['sol'].y[:, k], r4['sol'].t[k], frame='current')[1]
                       for k in range(len(r4['t']))])
        surf = np.array([m4.S.interp(r4['sol'].y[m4.ndof:, k], [1.0])[0]
                         for k in range(len(r4['t']))])       # sigma=1 处的真实表面值
        P4 = np.column_stack([P4, surf])
        Rt = np.array([float(rad(tt)) * 100.0 for tt in r4['t']])
        write_xlsx(os.path.join(OUT, 'result4.xlsx'), r4['t'], {'水分浓度': P4},
                   surface_extra=True,
                   extra_sheet={'表面位置': (['时间\\实际半径', '实际半径/cm'],
                                          list(zip(r4['t'], np.round(Rt, 4))))})
        # (b) 材料(Lagrangian)坐标版本：每列恒为同一材料点，末列恒为表面
        write_xlsx(os.path.join(OUT, 'result4_material_frame.xlsx'), r4['t'],
                   {'水分浓度': r4['C']}, surface_last=True)
        res['q4'] = dict(tf_s=r4['tf'], tf_h=r4['tf'] / 3600.0, table=r4['table'],
                         mono=r4['mono'])
    with open(os.path.join(OUT, 'summary_B.json'), 'w', encoding='utf-8') as f:
        json.dump(res, f, ensure_ascii=False, indent=1)
    print('done')
