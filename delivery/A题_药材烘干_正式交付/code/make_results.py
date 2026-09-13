# -*- coding: utf-8 -*-
# ==============================================================================
# make_results.py —— 生成四个问题的结果文件 result1~4.xlsx
# ==============================================================================
# 做什么：把谱配置法的解按题目模板采样写盘：第 1 行是到药材中心的距离，A 列是
#         时间，表体是各时刻各半径的温度或干基含水率，数值一律保留 4 位小数。
#         求解与写盘在同一趟里完成，不读已有的结果文件。
# 输入：spectral.HerbSolver 现场求解（物性、烘房环境与 R(t) 取 common.py，附件由它定位）。
#       生产网格 N=64，与 run_all.py 报的 t_dry、论文各表同一口径。
# 输出：<包根>/result/result1.xlsx ~ result4.xlsx。result1、result2 各含"温度""水分浓度"
#       两张表，result3、result4 各一张表；result4 的列是当时的物理距离，r_j > R(t_i) 的
#       位置已不存在故留空，末列"药材表面"给出该时刻的表面值 C(R(t), t)。
# 运行方式与依赖顺序：
#       python make_results.py    （可在任意目录运行：输出路径按本文件位置解析）
#       它不依赖 run_all.py 或 verify_data.py，但会重新求解四个问题（需要数分钟），
#       且重跑直接覆盖 result/*.xlsx，故没有必要时不要重跑。
# 关键变量：
#   N       谱节点数 64（论文的生产网格）
#   RCOLS   结果表的 21 个距离列 0.0~2.0，步长 0.1            [cm]
#   RC      同一组距离换算成米，用于插值                        [m]
#   CRIT    烘干判据阈值 0.15（决定 result3/4 的时间终点）      [kg/kg]
# ==============================================================================

# ---------------- 误 import 的防护 ----------------
# 本脚本没有 main()，生成逻辑直接写在模块顶层，import 即整段执行并覆盖 result/*.xlsx，
# 故拦一道：误 import 立刻报错，而不是悄悄覆盖掉已经交付的结果文件。
if __name__ != "__main__":
    raise RuntimeError(
        "make_results 的生成逻辑在模块顶层，import 会直接覆盖 result/*.xlsx。"
        "请用 `python make_results.py` 运行；确实需要 import 时，请先把生成逻辑移入 main()。")

import os
from console_utf8 import fix_console
fix_console()
import numpy as np
import openpyxl
import common as cm
from spectral import HerbSolver

OUTDIR = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                      "..", "result"))
os.makedirs(OUTDIR, exist_ok=True)

RCOLS = np.round(np.arange(0.0, 2.0001, 0.1), 10)      # 0.0 .. 2.0 cm
RC = RCOLS * 1e-2                                       # m
CRIT = 0.15
N = 64


def write_sheet(ws, times, values, header_last=None):
    """把一个（时刻 × 半径）的数值矩阵写成一张工作表。

    第 1 行是表头：A1 为模板规定的交叉标题，B1 起依次是 21 个距离 [cm]；
    A 列是各报告时刻 [s]；表体是 values[i, j]，四舍五入到 4 位小数。
    values 为 None 或非有限值时该单元格留空（问题 4 里已收缩掉的半径即如此处理）；
    header_last 不为 None 时，另用它在最后一列加一列表头与数值，如 result4 的"药材表面"。
    """
    ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
    for j, rc in enumerate(RCOLS):
        ws.cell(row=1, column=2 + j, value=float(rc))
    if header_last is not None:
        ws.cell(row=1, column=2 + len(RCOLS), value=header_last)
    for i in range(len(times)):
        ws.cell(row=2 + i, column=1, value=float(times[i]))
        for j in range(len(RCOLS)):
            v = values[i, j]
            if v is not None and np.isfinite(v):
                ws.cell(row=2 + i, column=2 + j, value=round(float(v), 4))
    if header_last is not None:
        for i in range(len(times)):
            v = values[i, len(RCOLS)]
            if v is not None and np.isfinite(v):
                ws.cell(row=2 + i, column=2 + len(RCOLS), value=round(float(v), 4))


def sample(s, sol, u_targets, which):
    """把解采样成矩阵，返回形状 (报告时刻数, 采样点数)。

    s         求解器实例（用它的 interp_u 做插值）；
    sol       solve_ivp 的解对象，sol.y 的各列依次对应各报告时刻；
    u_targets 目标归一化半径 u = (r/R)^2（无量纲，故与半径用什么单位无关）；
    which     "T" 取温度 [degC]，取其它值则取干基含水率 [kg/kg]。
    """
    out = np.empty((sol.y.shape[1], len(u_targets)))
    for i in range(sol.y.shape[1]):
        C, T = s.interp_u(sol.y[:, i], u_targets)
        out[i] = T if which == "T" else C
    return out


# ---------------- result1：问题 1（0~1800 s，步长 1 s） ----------------
# 温度与水分浓度各写一张表。A 列从 t=1 s 起（不含 t=0 的初始条件行），共 1800 行；
# 插值坐标取 u=(r/R)^2（谱配置法的自变量），故距离要先除以 R0 再平方。
print("result1 ...")
s = HerbSolver(N, prob=1)
tt = np.arange(1.0, 1801.0, 1.0)
at = np.concatenate([np.full(s.Np, 1e-13), np.full(s.Np, 1e-11)])
sol = s.solve(1800.0, t_eval=tt, rtol=1e-11, atol=at)
U = (RC / cm.R0) ** 2
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
write_sheet(ws, tt, sample(s, sol, U, "T"))
ws2 = wb.create_sheet("水分浓度")
write_sheet(ws2, tt, sample(s, sol, U, "C"))
wb.save(os.path.join(OUTDIR, "result1.xlsx"))

# ---------------- result2：问题 2（0~10800 s，步长 1 s） ----------------
# 结构与 result1 相同，只是物性换成随 C、T 变化的附录 3 关系，故需要解到 3 h。
print("result2 ...")
s2 = HerbSolver(N, prob=2)
tt2 = np.arange(1.0, 10801.0, 1.0)
at2 = np.concatenate([np.full(s2.Np, 1e-12), np.full(s2.Np, 1e-10)])
sol2 = s2.solve(10800.0, t_eval=tt2, rtol=1e-11, atol=at2)
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "温度"
write_sheet(ws, tt2, sample(s2, sol2, U, "T"))
ws2 = wb.create_sheet("水分浓度")
write_sheet(ws2, tt2, sample(s2, sol2, U, "C"))
wb.save(os.path.join(OUTDIR, "result2.xlsx"))


def dry(s, shrink=False, rtol=1e-11):
    """求烘干时长 t_dry [s]，事件口径与 run_all.dry_time 相同。

    事件函数取状态中水分那段的最大值减 CRIT：含水率沿半径由中心向外递减，最大值在轴心
    u=0（表面先干、中心最后达标），故它穿过 0 的时刻就是"处处不高于 0.15 kg/kg"的首达
    时刻；terminal 配 direction=-1 表示只在下降方向触发。rtol 是积分的相对容差。
    shrink 在本函数体内并不使用：是否含收缩由构造求解器 s 时决定，保留它只是让调用处
    一眼看出这一路是问题 4。
    """
    at = np.concatenate([np.full(s.Np, 1e-13), np.full(s.Np, 1e-11)])

    def ev(t, y):
        return np.max(y[:s.Np]) - CRIT
    ev.terminal = True
    ev.direction = -1
    e = s.solve(300000.0, rtol=rtol, atol=at, events=ev)
    return float(e.t_events[0][0])


# ---------------- result3：问题 3（到 t_dry，步长 60 s） ----------------
# 先由事件法解出终止时刻 t_dry，再把 60 s 整数倍的时刻连同 t_dry 本身作为输出行——
# 末行不是 60 s 的整数倍，它正是烘干刚好达标的那一刻。
print("result3 ...")
s3 = HerbSolver(N, prob=3)
td3 = dry(s3)
tt3 = np.arange(60.0, np.floor(td3 / 60.0) * 60.0 + 1.0, 60.0)
tt3 = np.append(tt3, td3)
at3 = np.concatenate([np.full(s3.Np, 1e-13), np.full(s3.Np, 1e-11)])
sol3 = s3.solve(td3, t_eval=list(tt3), rtol=1e-11, atol=at3)
U3 = (RC / cm.R0) ** 2
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
write_sheet(ws, tt3, sample(s3, sol3, U3, "C"))
wb.save(os.path.join(OUTDIR, "result3.xlsx"))

# ---------------- result4：问题 4（含收缩，到 t_dry，步长 60 s） ----------------
# 列的含义与 result3 相同，但距离列是"当时的物理距离"：每行按该时刻的 R(t) 判断哪些
# 距离仍然存在（r <= R(t) 才有值），不存在的留空，末列另给当前表面值。
print("result4 ...")
s4 = HerbSolver(N, prob=4, shrink=True)
td4 = dry(s4, shrink=True)
tt4 = np.arange(60.0, np.floor(td4 / 60.0) * 60.0 + 1.0, 60.0)
tt4 = np.append(tt4, td4)
at4 = np.concatenate([np.full(s4.Np, 1e-13), np.full(s4.Np, 1e-11)])
sol4 = s4.solve(td4, t_eval=list(tt4), rtol=1e-11, atol=at4)
vals = np.full((len(tt4), len(RCOLS) + 1), np.nan)
for i, t in enumerate(tt4):
    Rt = s4.R_of(t)
    for j, r in enumerate(RC):
        if r <= Rt + 1e-12:                 # 该物理位置仍然存在
            Cn, _ = s4.interp_u(sol4.y[:, i], (r / Rt) ** 2)
            vals[i, j] = Cn[0]
    vals[i, -1] = sol4.y[s4.N, i]           # 当前表面值
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Sheet1"
write_sheet(ws, tt4, vals, header_last="药材表面")
wb.save(os.path.join(OUTDIR, "result4.xlsx"))

print("result1: %d 行 x %d 列" % (len(tt), len(RCOLS) + 1))
print("result2: %d 行 x %d 列" % (len(tt2), len(RCOLS) + 1))
print("result3: %d 行, tdry=%.3f s" % (len(tt3), td3))
print("result4: %d 行, tdry=%.3f s" % (len(tt4), td4))
print("已写入:", OUTDIR)

